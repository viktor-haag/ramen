from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import (
    Appointment, AppointmentKind, AppStatus, FitRating, FitStatus, Posting, PostingFit, StatusChange,
)
from app.services import applications as svc
from app.services import export, triage

NOW = datetime(2026, 10, 25, 12, 0, tzinfo=UTC)  # 13:00 Berlin (CET after DST end)


@pytest.fixture
def db(app, client):  # client runs the lifespan (migrations)
    with app.state.session_factory() as s:
        yield s


def make_app(db, ext="1", status=AppStatus.saved, created=NOW - timedelta(days=5), **kw):
    p = Posting(source="src", external_id=ext, first_seen_at=NOW, last_seen_at=NOW,
                title=kw.pop("title", f"Title {ext}"), company=kw.pop("company", f"Co {ext}"),
                location=kw.pop("location", None), url=kw.pop("url", None))
    db.add(p)
    db.commit()
    a = triage.save_posting(db, p.id, created)
    a.status = status
    for k, v in kw.items():
        setattr(a, k, v)
    db.commit()
    return a


def change(db, a, to, at):
    db.add(StatusChange(application_id=a.id, from_status=a.status, to_status=to, changed_at=at))
    a.status = to
    a.updated_at = at
    db.commit()
    db.expire_all()  # session does not expire on commit; reload the collections


def test_rows_match_list_filter(db):
    a1 = make_app(db, "1", title="Python Dev")
    a2 = make_app(db, "2", title="Java Dev")
    a3 = make_app(db, "3", status=AppStatus.rejected, title="Python Old")
    a1.updated_at = NOW - timedelta(days=1)
    a2.updated_at = NOW
    db.commit()
    for statuses, q in [(None, None), ([AppStatus.rejected], None), (None, "python"),
                        ([AppStatus.saved, AppStatus.rejected], None)]:
        expected = [a.posting.title for a in svc.list_applications(db, statuses, q)]
        rows = export.export_rows(db, statuses, q, NOW, "de")
        assert [r["title"] for r in rows] == expected
    assert [r["title"] for r in export.export_rows(db, None, None, NOW, "de")] == ["Java Dev", "Python Dev"]
    assert a3.id  # rejected excluded by default


def test_applied_at_is_first_applied_change(db):
    a = make_app(db)
    assert export.export_rows(db, None, None, NOW, "de")[0]["applied_at"] is None
    t1, t2 = NOW - timedelta(days=3), NOW - timedelta(days=1)
    change(db, a, AppStatus.applied, t1)
    change(db, a, AppStatus.saved, t2 - timedelta(hours=1))
    change(db, a, AppStatus.applied, t2)
    row = export.export_rows(db, None, None, NOW, "de")[0]
    assert row["applied_at"] == t1
    assert row["last_change_at"] == t2
    assert row["saved_at"] == NOW - timedelta(days=5)


def test_next_appointment_is_earliest_future(db):
    a = make_app(db)
    db.add_all([
        Appointment(application_id=a.id, kind=AppointmentKind.deadline, starts_at=NOW - timedelta(days=1)),
        Appointment(application_id=a.id, kind=AppointmentKind.other, starts_at=NOW + timedelta(days=5)),
        Appointment(application_id=a.id, kind=AppointmentKind.interview,
                    starts_at=datetime(2026, 10, 27, 13, 30, tzinfo=UTC)),
    ])
    db.commit()
    assert export.export_rows(db, None, None, NOW, "de")[0]["next_appointment"] == "27.10.2026 14:30 Gespräch"
    assert export.export_rows(db, None, None, NOW, "en")[0]["next_appointment"] == "2026-10-27 14:30 Interview"


def test_none_next_appointment_and_fit_labels(db):
    a = make_app(db)
    row = export.export_rows(db, None, None, NOW, "de")[0]
    assert row["next_appointment"] is None and row["fit"] is None and row["fit_explanation"] is None
    db.add(PostingFit(posting_id=a.posting_id, status=FitStatus.done, rating=FitRating.moderate,
                      explanation="Teils."))
    db.commit()
    row = export.export_rows(db, None, None, NOW, "en")[0]
    assert row["fit"] == "Partial fit" and row["fit_explanation"] == "Teils."
    assert export.export_rows(db, None, None, NOW, "de")[0]["fit"] == "Passt teilweise"


def test_csv_format(db):
    a = make_app(db, title="A; B", company="Co", notes='line1\nline2')
    change(db, a, AppStatus.applied, datetime(2026, 10, 25, 13, 30, tzinfo=UTC))
    data = export.to_csv(export.export_rows(db, None, None, NOW, "de"), "de")
    assert data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    assert text.endswith("\r\n")
    header, body = text.split("\r\n", 1)
    assert header.split(";") == list(export.HEADERS["de"][c] for c in export.COLUMNS)
    assert '"A; B"' in body and '"line1\nline2"' in body
    assert "25.10.2026 14:30" in body and ";beworben;" in body
    en = export.to_csv(export.export_rows(db, None, None, NOW, "en"), "en").decode("utf-8-sig")
    assert "2026-10-25 14:30" in en and ";Applied;" in en


def test_csv_formula_guard(db):
    make_app(db, title='=HYPERLINK("x")', company="+1", location="@x", url="-y")
    text = export.to_csv(export.export_rows(db, None, None, NOW, "de"), "de").decode("utf-8-sig")
    assert '"\'=HYPERLINK(""x"")"' in text
    assert ";'+1;'@x;" in text
    assert ";'-y;" in text
    make_app(db, "2", title="\t=x", company="\rfoo")
    text = export.to_csv(export.export_rows(db, None, None, NOW, "de"), "de").decode("utf-8-sig")
    assert "'\t=x" in text and "\"'\rfoo\"" in text


def test_xlsx_structure(db):
    a = make_app(db, title="Dev", url="https://example.com/job")
    change(db, a, AppStatus.applied, datetime(2026, 10, 25, 13, 30, tzinfo=UTC))
    wb = load_workbook(BytesIO(export.to_xlsx(export.export_rows(db, None, None, NOW, "de"), "de")))
    ws = wb.active
    assert ws.title == "Bewerbungen"
    assert [c.value for c in ws[1]] == [export.HEADERS["de"][c] for c in export.COLUMNS]
    assert all(c.font.bold for c in ws[1])
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:O2"
    col = {c: i + 1 for i, c in enumerate(export.COLUMNS)}
    assert ws.cell(2, col["applied_at"]).value == datetime(2026, 10, 25, 14, 30)
    assert isinstance(ws.cell(2, col["saved_at"]).value, datetime)
    assert ws.cell(2, col["url"]).hyperlink.target == "https://example.com/job"
    assert ws.cell(2, col["status"]).value == "beworben"
    wb_en = load_workbook(BytesIO(export.to_xlsx([], "en")))
    assert wb_en.active.title == "Applications"


def test_xlsx_cells_are_strings(db):
    make_app(db, title='=SUM(1,2)', company="+x", notes="@y")
    ws = load_workbook(BytesIO(export.to_xlsx(export.export_rows(db, None, None, NOW, "de"), "de"))).active
    for c in (1, 2, export.COLUMNS.index("notes") + 1):
        assert ws.cell(2, c).data_type == "s"
    assert ws.cell(2, 1).value == "=SUM(1,2)"


@pytest.mark.parametrize("fmt,mt", [
    ("csv", "text/csv; charset=utf-8"),
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
])
def test_api_download_headers(client, db, monkeypatch, fmt, mt):
    monkeypatch.setattr("app.api.applications.utcnow", lambda: datetime(2026, 10, 25, 23, 30, tzinfo=UTC))
    r = client.get(f"/api/applications/export?format={fmt}")  # empty selection
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == mt
    # 23:30 UTC is already the 26th in Berlin
    assert r.headers["content-disposition"] == f'attachment; filename="ramen-bewerbungen-2026-10-26.{fmt}"'
    if fmt == "csv":
        assert r.content.decode("utf-8-sig").count("\r\n") == 1
    else:
        assert load_workbook(BytesIO(r.content)).active.max_row == 1
    en = client.get(f"/api/applications/export?format={fmt}&lang=en")
    assert 'filename="ramen-applications-2026-10-26.' in en.headers["content-disposition"]


def test_api_filters_and_validation(client, db):
    make_app(db, "1", title="Python Dev")
    make_app(db, "2", title="Java Dev")
    r = client.get("/api/applications/export?format=csv&q=python")
    assert r.content.decode("utf-8-sig").count("\r\n") == 2
    assert client.get("/api/applications/export?format=pdf").status_code == 422
    assert client.get("/api/applications/export?format=csv&lang=fr").status_code == 422


def test_overlong_cells_are_capped(db):
    make_app(db, title="T", notes="x" * 40_000, url="https://e.com/" + "a" * 2_500)
    rows = export.export_rows(db, None, None, NOW, "de")
    limit = export.MAX_CELL_CHARS
    ws = load_workbook(BytesIO(export.to_xlsx(rows, "de"))).active
    col = {c: i + 1 for i, c in enumerate(export.COLUMNS)}
    note = ws.cell(2, col["notes"]).value
    assert len(note) == limit and note.endswith("…")
    url_cell = ws.cell(2, col["url"])
    assert url_cell.hyperlink is None and url_cell.value.startswith("https://e.com/")
    import csv as _csv
    text = export.to_csv(rows, "de").decode("utf-8-sig")
    parsed = list(_csv.reader(text.split("\r\n")[:-1], delimiter=";"))
    cell = parsed[1][export.COLUMNS.index("notes")]
    assert len(cell) == limit and cell.endswith("…")
