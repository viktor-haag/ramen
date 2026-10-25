"""CSV / XLSX export of the application list (same filter and order as GET /applications)."""

import csv
import io
from datetime import datetime
from typing import Literal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import AppStatus, PostingFit
from app.services import applications as svc
from app.services.usage import BERLIN

Lang = Literal["de", "en"]

COLUMNS: tuple[str, ...] = (
    "title", "company", "location", "status", "source", "url", "saved_at", "applied_at",
    "last_change_at", "next_appointment", "salary_expectation", "channel", "fit",
    "fit_explanation", "notes",
)
DATE_COLUMNS = frozenset({"saved_at", "applied_at", "last_change_at"})

HEADERS: dict[str, dict[str, str]] = {
    "de": {
        "title": "Titel", "company": "Firma", "location": "Ort", "status": "Status",
        "source": "Quelle", "url": "Link", "saved_at": "Gemerkt am", "applied_at": "Beworben am",
        "last_change_at": "Letzte Statusänderung", "next_appointment": "Nächster Termin",
        "salary_expectation": "Gehaltsvorstellung", "channel": "Bewerbungsweg", "fit": "Fit",
        "fit_explanation": "Fit-Begründung", "notes": "Notizen",
    },
    "en": {
        "title": "Title", "company": "Company", "location": "Location", "status": "Status",
        "source": "Source", "url": "Link", "saved_at": "Saved on", "applied_at": "Applied on",
        "last_change_at": "Last status change", "next_appointment": "Next appointment",
        "salary_expectation": "Salary expectation", "channel": "Application channel",
        "fit": "Fit", "fit_explanation": "Fit explanation", "notes": "Notes",
    },
}

# Wording copied from frontend/src/i18n/locales/{de,en}.json (status.*, fit.rating.*, appointmentKind.*)
STATUS_LABELS = {
    "de": {"saved": "gemerkt", "preparing": "in Vorbereitung", "applied": "beworben",
           "interviewing": "im Gespräch", "offer": "Angebot", "rejected": "abgesagt",
           "withdrawn": "zurückgezogen", "accepted": "angenommen"},
    "en": {"saved": "Saved", "preparing": "Preparing", "applied": "Applied",
           "interviewing": "Interviewing", "offer": "Offer", "rejected": "Rejected",
           "withdrawn": "Withdrawn", "accepted": "Accepted"},
}
FIT_LABELS = {
    "de": {"good": "Passt gut", "moderate": "Passt teilweise", "poor": "Passt schlecht"},
    "en": {"good": "Good fit", "moderate": "Partial fit", "poor": "Poor fit"},
}
KIND_LABELS = {
    "de": {"deadline": "Frist", "interview": "Gespräch", "other": "Sonstiges"},
    "en": {"deadline": "Deadline", "interview": "Interview", "other": "Other"},
}
SHEET_TITLES = {"de": "Bewerbungen", "en": "Applications"}
DATE_FORMATS = {"de": "%d.%m.%Y %H:%M", "en": "%Y-%m-%d %H:%M"}
XLSX_DATE_FORMATS = {"de": "DD.MM.YYYY HH:MM", "en": "YYYY-MM-DD HH:MM"}
WIDTHS = {
    "title": 40, "company": 28, "location": 22, "status": 16, "source": 16, "url": 36,
    "saved_at": 18, "applied_at": 18, "last_change_at": 22, "next_appointment": 28,
    "salary_expectation": 20, "channel": 20, "fit": 16, "fit_explanation": 50, "notes": 60,
}


def _fmt(value: datetime, lang: str) -> str:
    return value.astimezone(BERLIN).strftime(DATE_FORMATS[lang])


def export_rows(
    session: Session, statuses: list[AppStatus] | None, q: str | None, now: datetime, lang: str
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for app in svc.list_applications(session, statuses, q):
        posting = app.posting
        fit = session.get(PostingFit, posting.id)
        applied = [c.changed_at for c in app.status_changes if c.to_status == AppStatus.applied]
        upcoming = [a for a in app.appointments if a.starts_at > now]
        nxt = min(upcoming, key=lambda a: a.starts_at) if upcoming else None
        rows.append({
            "title": posting.title,
            "company": posting.company,
            "location": posting.location,
            "status": STATUS_LABELS[lang][app.status.value],
            "source": posting.source,
            "url": posting.url,
            "saved_at": app.created_at,
            "applied_at": min(applied) if applied else None,
            "last_change_at": max((c.changed_at for c in app.status_changes), default=None),
            "next_appointment": (
                f"{_fmt(nxt.starts_at, lang)} {KIND_LABELS[lang][nxt.kind.value]}" if nxt else None
            ),
            "salary_expectation": app.salary_expectation,
            "channel": app.channel,
            "fit": FIT_LABELS[lang][fit.rating.value] if fit and fit.rating else None,
            "fit_explanation": fit.explanation if fit else None,
            "notes": app.notes,
        })
    return rows


def _csv_cell(key: str, value: object, lang: str) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return _fmt(value, lang)
    text = str(value)
    if text[:1] in ("=", "+", "-", "@"):
        text = "'" + text  # formula-injection guard (postings are untrusted input)
    return text


def to_csv(rows: list[dict[str, object]], lang: str) -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerow([HEADERS[lang][c] for c in COLUMNS])
    for row in rows:
        writer.writerow([_csv_cell(c, row.get(c), lang) for c in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rows: list[dict[str, object]], lang: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLES[lang]
    ws.append([HEADERS[lang][c] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r, row in enumerate(rows, start=2):
        for i, key in enumerate(COLUMNS, start=1):
            value = row.get(key)
            if value is None:
                continue
            cell = ws.cell(r, i)
            if isinstance(value, datetime):
                cell.value = value.astimezone(BERLIN).replace(tzinfo=None)
                cell.number_format = XLSX_DATE_FORMATS[lang]
                continue
            text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
            cell.value = text
            cell.data_type = "s"  # never a formula, even if the text starts with "="
            if key == "url" and text.startswith(("http://", "https://")):
                cell.hyperlink = text
                cell.font = Font(color="0563C1", underline="single")
            elif key in ("notes", "fit_explanation"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, key in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[key]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
